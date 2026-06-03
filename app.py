import os
import pandas as pd
import streamlit as st
from openpyxl import load_workbook
from PIL import Image

EXCEL_FILE = "World Cup 2026 Sweepstake Tracker.xlsx"
ASSET_FOLDER = "Assets"
BANNER_FILE = "Page_Banner.png"
PRICE_PER_TEAM = 5

PRIZES = {
    "Tournament Winner": "£100",
    "Tournament Runner Up": "£50",
    "Earliest Goal": "£30",
    "Earliest Yellow Card": "£30",
    "Earliest Red Card": "£30",
}


st.set_page_config(
    page_title="World Cup 2026 Sweepstake",
    page_icon="🏆",
    layout="wide"
)


def clean_columns(df):
    df.columns = (
        df.columns
        .astype(str)
        .str.strip()
        .str.replace("\n", " ")
        .str.replace("  ", " ")
    )
    return df


def find_column(df, wanted):
    for column in df.columns:
        clean_name = (
            column.lower()
            .replace(" ", "")
            .replace("_", "")
            .replace("-", "")
        )
        if clean_name == wanted:
            return column

    st.error(f"Could not find column: {wanted}")
    st.write("Columns found:", list(df.columns))
    st.stop()


@st.cache_data
def load_team_table():
    if not os.path.exists(EXCEL_FILE):
        st.error(f"Could not find {EXCEL_FILE}")
        st.stop()

    df = pd.read_excel(EXCEL_FILE, engine="openpyxl")
    return clean_columns(df)


@st.cache_data
def load_event_table():
    events = {}

    workbook = load_workbook(EXCEL_FILE, data_only=True)
    sheet = workbook.active

    header_row = None
    category_col = None
    time_col = None
    team_col = None
    player_col = None

    for row in sheet.iter_rows():
        values = [cell.value for cell in row]
        cleaned = [
            str(value).strip().lower() if value is not None else ""
            for value in values
        ]

        if "category" in cleaned and "time" in cleaned and "team" in cleaned and "player" in cleaned:
            header_row = row[0].row
            category_col = cleaned.index("category") + 1
            time_col = cleaned.index("time") + 1
            team_col = cleaned.index("team") + 1
            player_col = cleaned.index("player") + 1
            break

    if header_row is None:
        return events

    for row_number in range(header_row + 1, sheet.max_row + 1):
        category = sheet.cell(row_number, category_col).value

        if category is None or str(category).strip() == "":
            continue

        events[str(category).strip()] = {
            "time": sheet.cell(row_number, time_col).value,
            "team": sheet.cell(row_number, team_col).value,
            "player": sheet.cell(row_number, player_col).value,
        }

    return events


def get_event(events, category):
    data = events.get(category, {})

    time = data.get("time", "")
    team = data.get("team", "")
    player = data.get("player", "")

    return (
        "" if time is None else str(time),
        "" if team is None else str(team),
        "" if player is None else str(player),
    )


def status_dot(owner):
    if pd.notna(owner) and str(owner).strip() != "":
        return "🟢 Taken"
    return "⚪ Available"


st.markdown(
    """
    <style>
    .stApp {
        background: linear-gradient(180deg, #f3f7fb 0%, #ffffff 100%);
    }

    .main-card {
        background: white;
        border-radius: 18px;
        padding: 24px;
        box-shadow: 0 8px 24px rgba(15, 35, 75, 0.08);
        border: 1px solid #d9e6f5;
        text-align: center;
    }

    .metric-title {
        font-size: 15px;
        color: #0a1f44;
        font-weight: 700;
        margin-top: 6px;
    }

    .metric-value-blue {
        font-size: 42px;
        color: #0066cc;
        font-weight: 900;
    }

    .metric-value-green {
        font-size: 42px;
        color: #0a9d4f;
        font-weight: 900;
    }

    .metric-value-gold {
        font-size: 42px;
        color: #f2a900;
        font-weight: 900;
    }

    .section-card {
        background: white;
        border-radius: 18px;
        padding: 22px;
        box-shadow: 0 8px 24px rgba(15, 35, 75, 0.08);
        border: 1px solid #d9e6f5;
        margin-bottom: 18px;
    }

    .event-title {
        color: #0a1f44;
        font-size: 14px;
        font-weight: 900;
        text-align: center;
    }

    .event-main {
        font-size: 36px;
        font-weight: 900;
        text-align: center;
    }

    .event-sub {
        color: #0a1f44;
        font-size: 14px;
        text-align: center;
    }

    .taken-badge {
        background: #eef5fc;
        border-radius: 12px;
        padding: 14px;
        text-align: center;
        color: #0a1f44;
        font-weight: 900;
        font-size: 20px;
    }

    .footer-card {
        background: linear-gradient(90deg, #eef5fc, #ffffff);
        border-radius: 18px;
        padding: 24px;
        border: 1px solid #d9e6f5;
        color: #0a1f44;
    }

    [data-testid="stDataFrame"] {
        border-radius: 18px;
        overflow: hidden;
    }
    </style>
    """,
    unsafe_allow_html=True
)


df = load_team_table()
events = load_event_table()

nation_col = find_column(df, "nation")
owner_col = find_column(df, "ownedby")
status_col = find_column(df, "status")

total_teams = len(df)
taken_teams = df[owner_col].apply(
    lambda x: pd.notna(x) and str(x).strip() != ""
).sum()
prize_fund = total_teams * PRICE_PER_TEAM


banner_path = os.path.join(ASSET_FOLDER, BANNER_FILE)

if os.path.exists(banner_path):
    banner = Image.open(banner_path)
    st.image(banner, use_container_width=True)
else:
    st.title("🏆 World Cup 2026 Sweepstake")


col1, col2, col3, col4 = st.columns([1, 1, 1, 1.45])

with col1:
    st.markdown(
        f"""
        <div class="main-card">
            <div style="font-size: 48px;">⚽</div>
            <div class="metric-value-blue">{total_teams}</div>
            <div class="metric-title">Teams Available</div>
        </div>
        """,
        unsafe_allow_html=True
    )

with col2:
    st.markdown(
        f"""
        <div class="main-card">
            <div style="font-size: 48px;">💷</div>
            <div class="metric-value-green">£{PRICE_PER_TEAM}</div>
            <div class="metric-title">Per Team</div>
        </div>
        """,
        unsafe_allow_html=True
    )

with col3:
    st.markdown(
        f"""
        <div class="main-card">
            <div style="font-size: 48px;">🏆</div>
            <div class="metric-value-gold">£{prize_fund}</div>
            <div class="metric-title">Total Prize Fund</div>
        </div>
        """,
        unsafe_allow_html=True
    )

with col4:
    prize_html = "<div class='section-card'><h4 style='text-align:center;color:#0a1f44;'>PRIZE BREAKDOWN</h4>"
    for prize, amount in PRIZES.items():
        prize_html += f"""
        <div style="display:flex;justify-content:space-between;border-bottom:1px solid #e5edf6;padding:7px 0;color:#0a1f44;">
            <span>{prize}</span>
            <strong>{amount}</strong>
        </div>
        """
    prize_html += "</div>"
    st.markdown(prize_html, unsafe_allow_html=True)


earliest_goal_time, earliest_goal_team, earliest_goal_player = get_event(events, "Earliest Goal")
yellow_time, yellow_team, yellow_player = get_event(events, "Earliest Yellow Card")
red_time, red_team, red_player = get_event(events, "Earliest Red Card")
favourite_time, favourite_team, favourite_player = get_event(events, "Tournament Favourite")

event_cols = st.columns(4)

event_items = [
    ("FASTEST GOAL", f"{earliest_goal_time}’", f"{earliest_goal_player} ({earliest_goal_team})", "#0a9d4f"),
    ("EARLIEST YELLOW CARD", f"{yellow_time}’", f"{yellow_player} ({yellow_team})", "#f2a900"),
    ("EARLIEST RED CARD", f"{red_time}’", f"{red_player} ({red_team})", "#e33b2e"),
    ("TOURNAMENT FAVOURITE", favourite_team, favourite_player, "#0066cc"),
]

for column, item in zip(event_cols, event_items):
    title, main, sub, colour = item

    with column:
        st.markdown(
            f"""
            <div class="section-card">
                <div class="event-title">{title}</div>
                <div class="event-main" style="color:{colour};">{main}</div>
                <div class="event-sub">{sub}</div>
            </div>
            """,
            unsafe_allow_html=True
        )


left, right = st.columns([3, 1])

with left:
    st.markdown("## 🏆 Team Selections")

with right:
    st.markdown(
        f"""
        <div class="taken-badge">
            <span style="color:#0a9d4f;">{taken_teams}</span> / {total_teams} Teams Taken
        </div>
        """,
        unsafe_allow_html=True
    )


display_df = df.copy()

display_df["Owned By Display"] = display_df[owner_col].apply(
    lambda x: str(x).strip() if pd.notna(x) and str(x).strip() != "" else "Available"
)

display_df["Status Display"] = display_df[owner_col].apply(status_dot)

display_df = display_df[[nation_col, "Owned By Display", "Status Display"]]
display_df.insert(0, "#", range(1, len(display_df) + 1))

display_df = display_df.rename(
    columns={
        nation_col: "Team",
        "Owned By Display": "Owned By",
        "Status Display": "Status"
    }
)

search = st.text_input("Search teams or owners", "")

filter_option = st.radio(
    "Filter",
    ["All", "Available", "Taken"],
    horizontal=True
)

filtered_df = display_df.copy()

if search:
    search_lower = search.lower()
    filtered_df = filtered_df[
        filtered_df["Team"].astype(str).str.lower().str.contains(search_lower)
        | filtered_df["Owned By"].astype(str).str.lower().str.contains(search_lower)
    ]

if filter_option == "Available":
    filtered_df = filtered_df[filtered_df["Owned By"] == "Available"]

if filter_option == "Taken":
    filtered_df = filtered_df[filtered_df["Owned By"] != "Available"]

st.dataframe(
    filtered_df,
    use_container_width=True,
    hide_index=True,
    height=650
)


st.markdown(
    """
    <div class="footer-card">
        <h3>ℹ How it works</h3>
        <p>
        Pick an available team for £5. Multiple entries allowed.
        Winners are decided by the tournament winner, runner up and earliest in game events.
        </p>
        <p style="font-weight:900;">🏆 Good luck and enjoy the tournament!</p>
    </div>
    """,
    unsafe_allow_html=True
)
